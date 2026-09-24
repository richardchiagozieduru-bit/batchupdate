import pandas as pd
import openpyxl

def inspect_file(path):
    print("INSIDE FILE:", path)
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        print("Sheets:", wb.sheetnames)
        ws = wb.active
        # Read first few rows
        rows = list(ws.iter_rows(max_row=5, values_only=True))
        print("First 5 rows:")
        for r in rows:
            print("  ", r)
    except Exception as e:
        print("Error:", e)

print("--- 1013_06072026_uvu1.xlsx ---")
inspect_file("media/uploads/1013_06072026_uvu1.xlsx")

print("\n--- Un-updated_Accounts_as_at_202605_for_UVUOMA_MICROFINANCE_BANK_LIMITED_REAL_2026.xlsx ---")
inspect_file("media/uploads/Un-updated_Accounts_as_at_202605_for_UVUOMA_MICROFINANCE_BANK_LIMITED_REAL_2026.xlsx")
